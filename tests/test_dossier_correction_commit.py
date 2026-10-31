"""Real DOCX derivative orchestration with explicit interruption/replay boundaries."""
import io
from hashlib import sha256
from types import SimpleNamespace
from uuid import uuid4

import pytest
from docx import Document
from fastapi import HTTPException

from apps.api.access_control import AccessRejected
from apps.api.dossier_journey import CorrectionCommitBody, DossierStore
from apps.api.signed_supabase_originals import PrivateStorageOutcomeUnknown


def fixture():
    document = Document()
    document.add_paragraph('See Stuk 1, p. 2.')
    out = io.BytesIO()
    document.save(out)
    raw = out.getvalue()
    command = CorrectionCommitBody(snapshot_hash='a'*64, original_sha256=sha256(raw).hexdigest(),
        reference_id='ref-0000', paragraph_index=0, before='p. 2', after='p. 3', request_id=str(uuid4()))
    return raw, command


class Revisions:
    def __init__(self):
        self.row = {'state': 'absent'}
        self.activations = 0

    def reserve(self, bearer, case, request, action, values):
        if self.row['state'] != 'absent':
            if self.row['command_sha256'] != values['command_sha256']:
                raise HTTPException(409, 'conflict')
            return dict(self.row)
        if action == 'status':
            return dict(self.row)
        self.row = dict(state='reserved', parent_version_id=values['parent_version_id'],
                        parent_snapshot_hash=values['snapshot_hash'], reference_id=values['reference_id'],
                        document_id='document', version_id='new', original_sha256=values['sha256'],
                        byte_length=values['byte_length'], command_sha256=values['command_sha256'])
        return dict(self.row)

    def activate(self, bearer, case, request, digest, size):
        self.activations += 1
        self.row['state'] = 'active'
        return dict(self.row)


class Library:
    def __init__(self, original, mode):
        self.original, self.mode, self.data = original, mode, None
        self.writes = 0

    def put_once(self, *, data, **params):
        self.writes += 1
        self.data = data
        if self.mode == 'unknown':
            raise PrivateStorageOutcomeUnknown('lost acknowledgement')

    def read(self, *, version_id, **params):
        if version_id == 'old':
            return self.original
        return b'corrupted' if self.mode == 'corrupt' else self.data


def store(mode='normal'):
    raw, command = fixture()
    service = object.__new__(DossierStore)
    service.revisions = Revisions()
    service.client = SimpleNamespace(library=Library(raw, mode))
    service.load = lambda *args: ({'snapshot_hash': 'a'*64,
                                  'findings': [{'reference_id': 'ref-0000'}]}, raw)
    return service, command


@pytest.mark.parametrize('mode', ['normal', 'unknown'])
def test_fresh_readback_required_before_activation_and_committed_retry(mode):
    service, command = store(mode)
    result = service.commit_correction('signed', 'case', 'old', command)
    assert result['committed'] is True and result['requires_recheck'] is True
    assert service.revisions.activations == 1
    assert Document(io.BytesIO(service.client.library.data)).paragraphs[0].text == 'See Stuk 1, p. 3.'
    replay = service.commit_correction('signed', 'case', 'old', command)
    assert replay['version_id'] == result['version_id']
    assert service.client.library.writes == 1


def test_corrupt_readback_never_activates_and_reserved_retry_recovers():
    service, command = store('corrupt')
    with pytest.raises(AccessRejected):
        service.commit_correction('signed', 'case', 'old', command)
    assert service.revisions.activations == 0 and service.revisions.row['state'] == 'reserved'
    service.client.library.mode = 'normal'
    assert service.commit_correction('signed', 'case', 'old', command)['committed'] is True


def test_changed_command_and_stale_snapshot_rejected():
    service, command = store()
    with pytest.raises(HTTPException) as stale:
        service.commit_correction('signed', 'case', 'old', command.model_copy(update={'snapshot_hash': 'b'*64}))
    assert stale.value.status_code == 409 and service.client.library.writes == 0
    service.commit_correction('signed', 'case', 'old', command)
    with pytest.raises(HTTPException) as conflict:
        service.commit_correction('signed', 'case', 'old', command.model_copy(update={'after': 'p. 4'}))
    assert conflict.value.status_code == 409


def test_signed_reservation_cannot_substitute_another_snapshot_or_reference():
    service, command = store('corrupt')
    with pytest.raises(AccessRejected):
        service.commit_correction('signed', 'case', 'old', command)
    service.revisions.row['parent_snapshot_hash'] = 'f'*64
    with pytest.raises(HTTPException) as conflict:
        service.commit_correction('signed', 'case', 'old', command)
    assert conflict.value.status_code == 409 and service.revisions.activations == 0
