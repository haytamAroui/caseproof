import io

from docx import Document

from domains.disputes.dossier_verification import build_findings
from infrastructure.parsers.docx_reference_discovery import discover_docx_references
from infrastructure.parsers.docx_text_regions import extract_docx_text_regions


def test_english_exhibit_is_disclosed_without_becoming_piece_or_absent_inventory():
    document = Document()
    document.add_paragraph('Exhibit A, PDF p. 1; Exhibit 1, PDF p. 1; Pièce 1, PDF p. 1.')
    output = io.BytesIO()
    document.save(output)
    raw = output.getvalue()
    discovery = discover_docx_references(raw)
    assert len(discovery.references) == 1 and len(discovery.unsupported_references) == 2
    for reference in discovery.unsupported_references:
        assert document.paragraphs[0].text[reference.start:reference.end] == reference.raw
    regions = {region.part: region.text for region in extract_docx_text_regions(raw).regions}
    sources = [dict(role='exhibit', version_id='v', document_id='d', sha256='a'*64, page_count=1)]
    confirmation = {'payload': {'inventory_complete': True, 'entries': [
        {'label': 'piece', 'number': '1', 'source_version_id': 'v', 'pages': []}]}}
    report = build_findings(discovery, regions, sources, confirmation)
    assert report['findings'][0]['identity'] == 'resolved'
    for finding in report['findings'][1:]:
        assert finding['identity'] == 'not_assessed'
        assert finding['identity_reason'] == 'UNSUPPORTED_REFERENCE_DIALECT'
        assert finding['source_id'] is None and finding['pdf_page'] is None
        assert finding['page_mapping'] == 'unmapped'
        assert finding['fidelity'] == 'not_assessed' and finding['support'] == 'not_assessed'
    assert report['coverage']['detected'] == 3
    assert report['coverage']['unsupported_references'] == 2
    assert not report['coverage']['fully_checked']


def test_canonical_loader_preserves_unsupported_states_in_frozen_report():
    from apps.api.dossier_journey import DossierStore

    document = Document()
    document.add_paragraph('Exhibit 1, PDF p. 1.')
    data = io.BytesIO()
    document.save(data)

    class FixtureStore(DossierStore):
        def __init__(self):
            self.receipt_key = b'synthetic-canonical-unsupported-fixture'

        def rpc(self, bearer, case, action, data=None):
            if action == 'context':
                return {'bundle_signature': 'a'*64, 'confirmation': {
                    'request_id': 'confirmed', 'payload': {'inventory_complete': True, 'entries': [
                        {'label': 'piece', 'number': '1', 'source_version_id': None, 'pages': []}]}}}
            return {}

        def sources(self, bearer, case, *, inspect_pdf=True):
            return ([{'role': 'draft', 'version_id': 'draft'}, {'role': 'inventory', 'version_id': 'inventory'}],
                    {'draft': data.getvalue(), 'inventory': 'Pièce 1'.encode()})

    report, _ = FixtureStore().load('synthetic-signed-boundary', 'case', 'draft')
    assert report['coverage']['detected'] == 1
    finding = report['findings'][0]
    assert finding['identity_reason'] == 'UNSUPPORTED_REFERENCE_DIALECT'
    assert finding['identity'] == 'not_assessed' and finding['technical_check'] == 'unresolved'
    assert finding['fidelity'] == 'not_assessed'


def test_ordinary_english_verb_does_not_create_an_exhibit_notice():
    from domains.disputes.reference_detection import detect_unsupported_references
    assert not detect_unsupported_references('They exhibit a pattern of behaviour.', region='body')
    assert not detect_unsupported_references('Exhibit Appendix describes the method.', region='body')
    assert len(detect_unsupported_references('See exhibit A and Exhibit 12.', region='body')) == 2
