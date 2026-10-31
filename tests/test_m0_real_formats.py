"""M0 real document formats are rendered bytes, not declared MIME labels."""
from __future__ import annotations

import io
import json

import pytest
from docx import Document
from pypdf import PdfReader

from scripts.m0_real_format_fixtures import generate, verify


@pytest.fixture(scope="module")
def generated(tmp_path_factory: pytest.TempPathFactory):
    path = tmp_path_factory.mktemp("m0-real-formats")
    manifest = generate(path)
    return path, manifest


def test_all_trilingual_real_formats_are_generated_and_roundtrip(generated):
    out, manifest = generated
    counts = verify(out)
    assert counts["docx"] >= 18
    assert counts["text_pdf"] >= 12
    assert counts["image_pdf"] >= 3
    assert counts["inventory"] >= 18
    assert manifest["review_status"] == "author_supplied_unadjudicated"
    assert manifest["legal_validation"] is False
    assert manifest["production_ocr_quality_tested"] is False
    assert "CC0-1.0" in manifest["license_declaration"]


@pytest.mark.parametrize("lang", ["nl", "fr", "en"])
def test_each_language_has_docx_text_pdf_and_actual_image_scan(generated, lang):
    out, manifest = generated
    cases = [c for c in manifest["cases"] if c["language"] == lang]
    assert cases
    assert all((out / c["id"] / "conclusions.docx").is_file() for c in cases)
    all_sources = [s for c in cases for s in c["sources"]]
    assert any(s["image_only"] for s in all_sources), lang
    assert any(not s["image_only"] for s in all_sources), lang


def test_docx_has_real_ooxml_body_and_original_citation(generated):
    out, manifest = generated
    case = next(c for c in manifest["cases"] if c["id"] == "NL_OFFSET_21")
    document = Document(out / case["id"] / "conclusions.docx")
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Stuk 21" in text
    assert "gedrukte p. 17" in text
    assert len(document.tables) == 1
    assert "FICTIONAL" in text


def test_real_text_pdf_has_three_distinct_numbered_pages(generated):
    out, manifest = generated
    case = next(c for c in manifest["cases"] if c["id"] == "NL_OFFSET_21")
    source = next(s for s in case["sources"] if not s["image_only"])
    pdf = PdfReader(str(out / source["path"]))
    assert len(pdf.pages) == 3
    assert "Pakket ontvangen op 4 mei." in pdf.pages[2].extract_text()
    assert "Printed page label: 17" in pdf.pages[2].extract_text()
    assert source["bundle_page"] == 42  # independently mapped, not source PDF pagination


def test_image_only_pdf_has_real_image_and_no_hidden_pdf_text(generated):
    out, manifest = generated
    case = next(c for c in manifest["cases"] if c["id"] == "EN_SCAN_37")
    source = next(s for s in case["sources"] if s["image_only"])
    pdf = PdfReader(str(out / source["path"]))
    assert len(pdf.pages) == 2
    for page in pdf.pages:
        assert not (page.extract_text() or "").strip()
        assert page.images


def test_corrupted_original_detected_by_frozen_manifest(generated, tmp_path):
    out, manifest = generated
    import shutil

    cloned = tmp_path / "tampered"
    shutil.copytree(out, cloned)
    record = manifest["file_records"][0]
    target = cloned / record["path"]
    target.write_bytes(target.read_bytes() + b"tamper")
    with pytest.raises(ValueError, match="fixture hash/size mismatch"):
        verify(cloned)


def test_manifest_allows_no_external_lawyer_signoff(generated):
    _, manifest = generated
    raw = json.dumps(manifest)
    assert '"legal_validation": false' in raw
    assert '"review_status": "author_supplied_unadjudicated"' in raw
    assert '"production_ocr_quality_tested": false' in raw


def test_scanned_pdf_is_not_a_claim_of_ocr_transcription(generated):
    out, manifest = generated
    scan = next(x for x in manifest["file_records"] if x["role"] == "image_pdf")
    raw = (out / scan["path"]).read_bytes()
    assert raw.startswith(b"%PDF-")
    assert all(not (page.extract_text() or "").strip() for page in PdfReader(io.BytesIO(raw)).pages)
