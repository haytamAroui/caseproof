import io
from hashlib import sha256
from zipfile import ZipFile

import pytest
from docx import Document

from infrastructure.parsers.dossier_docx_export import corrected_docx_preview


def draft(split=False):
    document = Document()
    paragraph = document.add_paragraph('See Stuk 1, p. 2.')
    if split:
        paragraph.add_run(' More text.').bold = True
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_correction_preserves_original_and_other_package_parts():
    original = draft()
    result = corrected_docx_preview(original, expected_sha256=sha256(original).hexdigest(),
                                    paragraph_index=0, before='p. 2', after='p. 3')
    assert Document(io.BytesIO(original)).paragraphs[0].text == 'See Stuk 1, p. 2.'
    assert Document(io.BytesIO(result)).paragraphs[0].text == 'See Stuk 1, p. 3.'
    with ZipFile(io.BytesIO(original)) as old, ZipFile(io.BytesIO(result)) as new:
        assert old.namelist() == new.namelist()
        assert b'xmlns:w14=' in new.read('word/document.xml')
        for name in old.namelist():
            if name != 'word/document.xml':
                assert old.read(name) == new.read(name)


@pytest.mark.parametrize('split,before,digest', [(True, 'p. 2', None), (False, 'missing', None),
                                               (False, 'p. 2', '0' * 64)])
def test_complex_stale_and_missing_ranges_abstain(split, before, digest):
    original = draft(split)
    with pytest.raises(ValueError):
        corrected_docx_preview(original, expected_sha256=digest or sha256(original).hexdigest(),
                               paragraph_index=0, before=before, after='p. 3')
