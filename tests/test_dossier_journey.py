"""Offline fail-closed acceptance; live signed proof is a separate workflow."""
from __future__ import annotations

import hashlib
import io
from uuid import uuid4
from zipfile import ZipFile

from domains.disputes.dossier_verification import build_findings
from infrastructure.parsers.docx_reference_discovery import discover_docx_references
from infrastructure.parsers.docx_text_regions import extract_docx_text_regions
from infrastructure.parsers.dossier_docx_export import annotated_docx
from tests.test_phase4_headless_api import docx


def assess(draft, sources, confirmation):
    return build_findings(discover_docx_references(draft),
                          {r.part: r.text for r in extract_docx_text_regions(draft).regions}, sources, confirmation)


def test_identity_needs_confirmation_and_absence_needs_complete_inventory():
    draft = docx('Stuk 1, PDF p. 1; Stuk 9, PDF p. 1.')
    source = {'role': 'exhibit', 'version_id': 'v', 'document_id': 'd', 'sha256': 'a' * 64, 'page_count': 1}
    assert [f['identity'] for f in assess(draft, [source], None)['findings']] == ['not_assessed'] * 2
    context = {'payload': {'inventory_complete': False, 'entries': [
        {'label': 'stuk', 'number': '1', 'source_version_id': 'v', 'pages': []}]}}
    report = assess(draft, [source], context)
    assert report['findings'][0]['technical_check'] == 'unresolved'
    assert report['findings'][0]['identity_reason'] == 'INVENTORY_INCOMPLETE'
    assert report['findings'][1]['identity'] == 'not_assessed'
    context['payload']['inventory_complete'] = True
    assert assess(draft, [source], context)['findings'][1]['identity'] == 'not_in_inventory'


def test_crops_do_not_create_page_mapping_and_conflicts_abstain():
    source = {'role': 'exhibit', 'version_id': 'v', 'document_id': 'd', 'sha256': 'a' * 64,
              'page_count': 2, 'geometry': {'1': {'crop': {'page_number': 1}}}, 'image_only': True}
    entry = {'label': 'piece', 'number': '2', 'source_version_id': 'v', 'pages': []}
    context = {'payload': {'inventory_complete': True, 'entries': [entry]}}
    assert assess(docx('Pièce 2, p. 1.'), [source], context)['findings'][0]['page_mapping'] == 'unmapped'
    entry['pages'] = [{'space': 'printed', 'label': '7', 'pdf_page': p, 'rationale': 'fixture'} for p in (1, 2)]
    finding = assess(docx('Pièce 2, printed p. 7.'), [source], context)['findings'][0]
    assert finding['page_mapping'] == 'conflicting' and finding['ocr_crop'] is None
    assert assess(docx('Pièce 2, PDF p. 99.'), [source], context)['coverage']['checked'] == 0


def test_empty_and_unreadable_never_fully_checked():
    for raw in (b'not a DOCX', docx('No explicit references.')):
        report = assess(raw, [], None)
        assert not report['coverage']['fully_checked']
        assert not report['coverage']['reference_recall_measured']


def test_annotated_derivative_preserves_original_and_reports_limitations():
    raw = docx('Zie Stuk 1.')
    digest = hashlib.sha256(raw).hexdigest()
    report = {**assess(raw, [], None), 'snapshot_hash': 'b' * 64, 'human_decisions': []}
    derivative = annotated_docx(raw, report)
    assert hashlib.sha256(raw).hexdigest() == digest and derivative != raw
    with ZipFile(io.BytesIO(derivative)) as archive:
        text = archive.read('word/document.xml').decode()
        assert 'scoped review annotations' in text and 'not_assessed' in text
        assert 'fully_checked' in text


def test_fixed_dossier_ground_truth():
    import pytest
    pytest.importorskip('docx')
    from tests.fixtures.dossier_journey import (
        EXPECTED_COUNTS,
        EXPECTED_IDENTITIES,
        confirmation,
        dossier_bundle,
    )
    bundle = dossier_bundle()
    text, scan = str(uuid4()), str(uuid4())
    sources = [{'role': 'exhibit', 'version_id': version, 'document_id': str(uuid4()),
                'sha256': source.sha256, 'page_count': 1} for source, version in zip(bundle[2:], (text, scan), strict=True)]
    c = confirmation('a' * 64, str(uuid4()), text, scan)
    report = assess(bundle[0].content, sources, {'payload': c})
    assert [f['identity'] for f in report['findings']] == EXPECTED_IDENTITIES
    assert {k: report['coverage'][k] for k in EXPECTED_COUNTS} == EXPECTED_COUNTS
    assert report['coverage']['unprocessed_regions'] >= 1


def test_omitted_confirmation_entry_cannot_become_absent_inventory():
    from apps.api.dossier_journey import DossierStore

    class FixtureStore(DossierStore):
        def __init__(self):
            self.receipt_key=b"synthetic-geometry-fixture-only"

        def rpc(self, bearer, case, action, data=None):
            if action == 'context':
                return {'bundle_signature': 'a' * 64, 'confirmation': {
                    'request_id': 'confirmed', 'payload': {'inventory_complete': True, 'entries': []}}}
            return {}

        def sources(self, bearer, case, *, inspect_pdf=True):
            raw = docx('Stuk 1, PDF p. 1.')
            return ([{'role': 'draft', 'version_id': 'draft'},
                     {'role': 'inventory', 'version_id': 'inventory'}],
                    {'draft': raw, 'inventory': b'Stuk 1 invoice'})

    report, _ = FixtureStore().load('signed-test-double', 'case', 'draft')
    assert report['findings'][0]['identity'] == 'not_assessed'
    assert report['findings'][0]['identity_reason'] == 'INVENTORY_INCOMPLETE'
    assert not report['coverage']['inventory_confirmed_complete']



def test_real_docx_export_preserves_namespace_declarations():
    import pytest
    pytest.importorskip('docx')
    from docx import Document
    original = Document()
    original.add_paragraph('Stuk 1, PDF p. 1.')
    data = io.BytesIO()
    original.save(data)
    raw = data.getvalue()
    report = {**assess(raw, [], None), 'snapshot_hash': 'c' * 64, 'human_decisions': []}
    derivative = annotated_docx(raw, report)
    assert 'scoped review annotations' in '\n'.join(p.text for p in Document(io.BytesIO(derivative)).paragraphs)
    with ZipFile(io.BytesIO(raw)) as old, ZipFile(io.BytesIO(derivative)) as new:
        assert b'xmlns:w14=' in old.read('word/document.xml')
        assert b'xmlns:w14=' in new.read('word/document.xml')



def test_fixture_pdfs_have_valid_structure_and_distinct_text_coverage():
    import pytest
    pytest.importorskip('docx')
    from pypdf import PdfReader

    from tests.fixtures.dossier_journey import dossier_bundle
    bundle = dossier_bundle()
    text = PdfReader(io.BytesIO(bundle[2].content), strict=True)
    scan = PdfReader(io.BytesIO(bundle[3].content), strict=True)
    assert len(text.pages) == len(scan.pages) == 1
    assert 'STUK 1 INVOICE' in text.pages[0].extract_text()
    assert not scan.pages[0].extract_text().strip()



def test_annotation_keeps_conflicting_human_decisions_visible():
    raw = docx('Stuk 1.')
    report = {**assess(raw, [], None), 'snapshot_hash': 'd' * 64,
              'human_decisions': [{'reference_id': 'ref-0000', 'decision': decision, 'rationale': 'synthetic'}
                                  for decision in ('CONFIRMED_SUPPORTED', 'CONFIRMED_CONTRADICTED')]}
    with ZipFile(io.BytesIO(annotated_docx(raw, report))) as archive:
        xml = archive.read('word/document.xml')
        assert b'CONFIRMED_SUPPORTED' in xml and b'CONFIRMED_CONTRADICTED' in xml


def test_page_sequences_never_silently_check_only_the_first_page():
    source = {'role': 'exhibit', 'version_id': 'v', 'document_id': 'd', 'sha256': 'a' * 64, 'page_count': 1}
    context = {'payload': {'inventory_complete': True, 'entries': [
        {'label': 'stuk', 'number': '1', 'source_version_id': 'v', 'pages': []}]}}
    for suffix in ('PDF p. 1-99', 'PDF pp. 1–99', 'PDF p. 1—99', 'PDF p. 1−99', 'PDF p. 1 à 99', 'PDF p. 1, 99',
                   'printed p. 1-99', 'bundle p. 1-99', 'p. 1-99'):
        result = assess(docx('Stuk 1, ' + suffix), [source], context)
        finding = result['findings'][0]
        assert finding['pdf_page'] is None and finding['technical_check'] == 'unresolved'
        assert finding['page_mapping_reason'] == 'page_sequence_requires_review'
        assert result['coverage']['checked'] == 0


def test_mixed_pdf_physical_reference_looks_up_its_actual_scanned_page():
    from types import SimpleNamespace

    from apps.api.access_control import AccessFailure, AccessRejected
    from apps.api.dossier_journey import DossierStore

    class GeometryClient:
        def __init__(self, available):
            self.available = available
            self.lookups = []

        def read_geometry(self, token, version, page):
            self.lookups.append((version, page))
            if not self.available:
                raise AccessRejected(AccessFailure.RESOURCE_MASKED)
            return (SimpleNamespace(original_sha256='a' * 64, geometry_sha256='b' * 64),
                    {'words': [{'left': 1, 'top': 2, 'width': 3, 'height': 4}]})

    class MixedStore(DossierStore):
        def __init__(self, available):
            self.receipt_key=b"synthetic-geometry-fixture-only"
            self.client = GeometryClient(available)
            from infrastructure.parsers.pdf_ocr import OCRUnavailable

            def unavailable(*args, **kwargs):
                raise OCRUnavailable("test_geometry_only")
            self.visual = SimpleNamespace(inspect=unavailable)

        def rpc(self, token, case, action, data=None):
            return {'bundle_signature': 'c' * 64, 'confirmation': {
                'request_id': 'revision', 'payload': {'inventory_complete': True, 'entries': [
                    {'label': 'piece', 'number': '2', 'source_version_id': 'mixed', 'pages': []}]}}}

        def sources(self, token, case, *, inspect_pdf=True):
            return ([{'role': 'draft', 'version_id': 'draft'}, {'role': 'inventory', 'version_id': 'inv'},
                     {'role': 'exhibit', 'version_id': 'mixed', 'document_id': 'doc', 'sha256': 'a' * 64,
                      'page_count': 2, 'image_only': False, 'image_only_pages': [2]}],
                    {'draft': docx('Pièce 2, PDF p. 2.'), 'inv': b'Piece 2 scan', 'mixed': b'geometry-only test double'})

    for available in (True, False):
        store = MixedStore(available)
        report, _ = store.load('signed-test-double', 'case', 'draft')
        finding = report['findings'][0]
        assert store.client.lookups == [('mixed', 2)]
        assert finding['pdf_page'] == 2 and finding['fidelity'] == 'ocr_unreliable'
        assert finding['ocr_uncertainty'] == ('requires_human_confirmation' if available else 'image_only_no_usable_ocr')
        assert bool(finding['ocr_crop']) == available
