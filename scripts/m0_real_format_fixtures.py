"""Generate and verify real-format, fictional NL/FR/EN/mixed M0 source fixtures.

Only source materials are synthetic. This does NOT adjudicate legal gold labels,
validate OCR accuracy or exercise the product end-to-end. Outputs are generated
from a pinned author-supplied v3 corpus and checked by independent file readers.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import textwrap
from pathlib import Path
from typing import Any

import reportlab
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from scripts.audit_m0_trilingual import DEFAULT_CORPUS, PRIOR_CORPUS, audit

FONT_PATH = Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"
PDF_FONT = "CaseProofM0Vera"
REVIEW_FLAGS = ("synthetic_only", "author_supplied_unadjudicated")


def _sha(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


def _safe(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]", "_", value)


def _pdf_font() -> None:
    if PDF_FONT not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(PDF_FONT, str(FONT_PATH)))


def _wrapped(text: str, limit: int = 67) -> list[str]:
    # Preserve non-ASCII and use a controlled line length for readable pages.
    return textwrap.wrap(text, width=limit, break_long_words=True) or ["(no text)"]


def _text_pdf(source: dict[str, Any], case_id: str) -> bytes:
    _pdf_font()
    out = io.BytesIO()
    pdf = canvas.Canvas(out, pagesize=(595, 842), invariant=1, pageCompression=1)
    for page_number in range(1, source["pdf_page"] + 1):
        pdf.setFont(PDF_FONT, 12)
        pdf.drawString(45, 786, "CaseProof M0 - fictional test material")
        pdf.drawString(45, 764, f"{case_id} / {source['source_id']} / PDF page {page_number}")
        if page_number == source["pdf_page"]:
            y = 710
            for line in _wrapped(source["text"]):
                pdf.drawString(45, y, line)
                y -= 22
        else:
            pdf.drawString(45, 708, "Synthetic blank preceding source page.")
        if page_number == source["pdf_page"] and source["printed_page"] is not None:
            pdf.drawString(45, 65, f"Printed page label: {source['printed_page']}")
        pdf.showPage()
    pdf.save()
    return out.getvalue()


def _scanned_pdf(source: dict[str, Any], case_id: str, language: str) -> bytes:
    """Real image-only PDF: every page is raster; PDF has NO text layer."""
    out = io.BytesIO()
    pdf = canvas.Canvas(out, pagesize=(595, 842), invariant=1, pageCompression=1)
    font = ImageFont.truetype(str(FONT_PATH), 31)
    small = ImageFont.truetype(str(FONT_PATH), 22)
    for page_number in range(1, source["pdf_page"] + 1):
        im = Image.new("RGB", (1190, 1684), "white")
        draw = ImageDraw.Draw(im)
        draw.text((95, 135), "CASEPROOF - SYNTHETIC SCAN", font=font, fill="black")
        draw.text((95, 205), f"{case_id} / {language.upper()}", font=small, fill="black")
        draw.text((95, 265), f"Original raster page {page_number}", font=small, fill="black")
        draw.text((95, 325), "No searchable text layer", font=small, fill="black")
        if page_number == source["pdf_page"]:
            draw.text((95, 465), "IMAGE-ONLY PAGE / OCR NOT PERFORMED", font=small, fill="black")
            if source.get("text"):
                for index, line in enumerate(_wrapped(source["text"], 45)):
                    draw.text((95, 530 + index * 42), line, font=small, fill="black")
            if source["printed_page"] is not None:
                draw.text((95, 1510), f"Printed label {source['printed_page']}", font=small, fill="black")
        pdf.drawImage(ImageReader(im), 0, 0, width=595, height=842)
        pdf.showPage()
    pdf.save()
    return out.getvalue()


def _draft_docx(case: dict[str, Any]) -> bytes:
    document = Document()
    document.core_properties.title = f"CaseProof synthetic draft {case['id']}"
    document.core_properties.author = "CaseProof synthetic fixture generator"
    document.add_heading("CaseProof synthetic conclusions", level=1)
    document.add_paragraph(f"Language cohort: {case['language'].upper()}")
    document.add_paragraph(case["brief_text"])
    document.add_heading("Source reference as supplied", level=2)
    document.add_paragraph(f"Exhibit label: {case['citation']['label']}")
    document.add_paragraph(f"PDF page: {case['citation']['pdf_page']}")
    if case["citation"]["quote"] is not None:
        document.add_paragraph("Quoted words: " + case["citation"]["quote"])
    else:
        document.add_paragraph("No direct quotation is asserted.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Fixture case ID"
    table.cell(0, 1).text = case["id"]
    document.add_paragraph("FICTIONAL - not a court filing or legal opinion.")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _register(output: Path, name: str, content: bytes, role: str, records: list[dict[str, Any]]) -> None:
    path = output / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)
    records.append({
        "path": name,
        "role": role,
        "bytes": len(content),
        "sha256": _sha(content),
    })


def generate(output: Path, corpus: Path = DEFAULT_CORPUS) -> dict[str, Any]:
    data = json.loads(corpus.read_text(encoding="utf-8"))
    predecessor = json.loads(PRIOR_CORPUS.read_text(encoding="utf-8"))
    audit(data, predecessor)
    output.mkdir(parents=True, exist_ok=True)
    records: list[dict[str, Any]] = []
    cases: list[dict[str, Any]] = []
    for case in data["cases"][data["legacy_case_count"]:]:
        cid = case["id"]
        prefix = f"{cid}/"
        _register(output, prefix + "conclusions.docx", _draft_docx(case), "draft_docx", records)
        inv = case["inventory"]
        inventory = (
            f"CASE={cid}\nSTATUS={inv['status']}\n"
            + "".join(f"ITEM={label}\n" for label in inv["listed_labels"])
        ).encode("utf-8")
        _register(output, prefix + "inventory.txt", inventory, "inventory", records)
        sources = []
        for index, source in enumerate(case["sources"], start=1):
            name = prefix + f"original-{index:02d}-{_safe(source['source_id'])}.pdf"
            scan = source["representation_status"] == "unprocessed_scan"
            content = (_scanned_pdf(source, cid, case["language"]) if scan else _text_pdf(source, cid))
            _register(output, name, content, "image_pdf" if scan else "text_pdf", records)
            sources.append({
                "path": name,
                "source_id": source["source_id"],
                "label": source["label"],
                "version": source["version"],
                "pdf_page": source["pdf_page"],
                "printed_page": source["printed_page"],
                "bundle_page": source["bundle_page"],
                "image_only": scan,
            })
        if cid == "FR_CLEAN_25":
            # A separately identified image-only alternate to an FR text-source
            # document; never treat it as a second preserved original for that
            # same exhibit or as an OCR-success claim.
            alternate = prefix + "alternate-image-only-control.pdf"
            _register(output, alternate, _scanned_pdf(case["sources"][0], cid, "fr"), "image_pdf", records)
            sources.append({
                "path": alternate,
                "source_id": case["sources"][0]["source_id"],
                "label": case["sources"][0]["label"],
                "version": case["sources"][0]["version"],
                "pdf_page": case["sources"][0]["pdf_page"],
                "printed_page": case["sources"][0]["printed_page"],
                "bundle_page": case["sources"][0]["bundle_page"],
                "image_only": True,
                "alternate_format_only": True,
            })
        cases.append({
            "id": cid,
            "language": case["language"],
            "inventory_status": inv["status"],
            "sources": sources,
            "author_expected_dimensions": case["expected_dimensions"],
            "not_human_adjudicated": True,
        })
    manifest: dict[str, Any] = {
        "schema_version": "caseproof-m0-real-formats-v1",
        "source_corpus_sha256": _sha(corpus.read_bytes()),
        "license_declaration": "CC0-1.0 project-authored synthetic text and derived fixture bytes",
        "material_provenance": "newly generated fictional samples, no real litigant/customer/court data",
        "review_status": "author_supplied_unadjudicated",
        "legal_validation": False,
        "production_ocr_quality_tested": False,
        "formats": ["DOCX", "text-layer PDF", "image-only scanned PDF", "UTF-8 inventory"],
        "file_records": records,
        "cases": cases,
    }
    (output / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return manifest


def verify(output: Path) -> dict[str, int]:
    manifest = json.loads((output / "manifest.json").read_text(encoding="utf-8"))
    if manifest["review_status"] != "author_supplied_unadjudicated" or manifest["legal_validation"]:
        raise ValueError("generator cannot claim independent review or legal validation")
    if manifest["source_corpus_sha256"] != _sha(DEFAULT_CORPUS.read_bytes()):
        raise ValueError("source corpus version changed")
    counts: dict[str, int] = {"docx": 0, "text_pdf": 0, "image_pdf": 0, "inventory": 0}
    for record in manifest["file_records"]:
        raw = (output / record["path"]).read_bytes()
        if _sha(raw) != record["sha256"] or len(raw) != record["bytes"]:
            raise ValueError(f"fixture hash/size mismatch: {record['path']}")
        role = record["role"]
        if role == "draft_docx":
            doc = Document(io.BytesIO(raw))
            text = "\n".join(p.text for p in doc.paragraphs)
            if "FICTIONAL" not in text or "CaseProof synthetic conclusions" not in text:
                raise ValueError("DOCX is missing actual expected text")
            counts["docx"] += 1
        elif role in {"text_pdf", "image_pdf"}:
            pdf = PdfReader(io.BytesIO(raw))
            if not pdf.pages:
                raise ValueError("empty PDF")
            extracted = "\n".join(p.extract_text() or "" for p in pdf.pages)
            if role == "image_pdf":
                if extracted.strip():
                    raise ValueError("image-only scan has a hidden/extracted text layer")
                if not all(page.images for page in pdf.pages):
                    raise ValueError("scanned PDF must contain actual raster imagery on every page")
            elif "CaseProof M0" not in extracted:
                raise ValueError("text PDF has no extractable source text")
            counts[role] += 1
        elif role == "inventory":
            if not raw.startswith(b"CASE=") or b"STATUS=" not in raw:
                raise ValueError("invalid inventory text")
            counts["inventory"] += 1
        else:
            raise ValueError(f"unexpected material role: {role}")
    if any(counts[k] < 3 for k in ("docx", "text_pdf", "image_pdf")):
        raise ValueError("missing tri-language real-format coverage")
    for case in manifest["cases"]:
        cid = case["id"]
        if not (output / cid / "conclusions.docx").is_file():
            raise ValueError(f"missing draft: {cid}")
        for source in case["sources"]:
            if len(PdfReader(str(output / source["path"])).pages) != source["pdf_page"]:
                raise ValueError(f"wrong physical PDF page count: {cid}")
    for language in ("nl", "fr", "en"):
        language_sources = [s for c in manifest["cases"] if c["language"] == language for s in c["sources"]]
        if not language_sources or not any(s["image_only"] for s in language_sources):
            raise ValueError(f"missing {language} image-only scanned PDF")
        if not any(not s["image_only"] for s in language_sources):
            raise ValueError(f"missing {language} searchable text PDF")
    return counts


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--verify-only", action="store_true")
    args = parser.parse_args()
    if not args.verify_only:
        generate(args.out)
    print(json.dumps({"status": "real_format_synthetic_only", "counts": verify(args.out)}, sort_keys=True))


if __name__ == "__main__":
    main()
