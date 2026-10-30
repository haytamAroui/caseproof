"""Generate deterministic 150-page / 60-reference synthetic M5 stress input.

This generator does not claim API benchmark results. Run against the signed
development API separately and report wall time, precision and recall.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path


def main() -> None:
    from docx import Document
    from docx.enum.text import WD_BREAK

    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=Path("m5-synthetic-150.docx"))
    parser.add_argument("--truth", type=Path, default=Path("m5-synthetic-truth.json"))
    args = parser.parse_args()
    document = Document()
    truth = []
    for page in range(1, 151):
        paragraph = document.add_paragraph(f"Synthetic dossier page {page}. ")
        if page <= 60:
            label = "Stuk" if page % 2 else "Pièce"
            number = page if page % 5 else page + 100
            page_number = page if page % 7 else 999
            paragraph.add_run(f"Zie {label} {number}, p. {page_number}.")
            truth.append({
                "synthetic_page": page, "label": label, "number": number,
                "requested_page": page_number,
                "scenario": "mismatched_number" if page % 5 == 0
                else "missing_page" if page % 7 == 0 else "reference_candidate",
                "identity_verified": False, "fidelity_verified": False,
                "semantic_support_verified": False,
            })
        if page != 150:
            paragraph.add_run().add_break(WD_BREAK.PAGE)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.truth.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    args.truth.write_text(json.dumps({
        "schema": "caseproof.m5.synthetic.v1", "intended_pages": 150,
        "references": truth, "note": "DOCX explicit page breaks, not a guarantee of renderer pagination; no exhibit scans are generated."
    }, indent=2) + "\n", encoding="utf-8")
    print(f"Generated {len(truth)} synthetic references and 150 intended pages")
    print("Benchmark: NOT RUN. Requires source bundle, OCR scans and signed API evaluation.")


if __name__ == "__main__":
    main()
