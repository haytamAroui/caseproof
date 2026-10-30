"""Real signed synthetic acceptance for the complete dossier API."""
from __future__ import annotations

import base64
import hashlib
import http.client
import io
import json
import os
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

from apps.api.supabase_password_signin import PasswordCredentials, SupabasePasswordSignin
from scripts.m05_registry import record
from scripts.phase3_live_acceptance import direct
from scripts.verify_dev_ephemeral_auth import PROJECT_REF
from scripts.verify_dev_live_review_roundtrip import _server_start, _server_stop
from scripts.verify_dev_live_txt_journey import admin_rpc, success
from tests.fixtures.dossier_journey import EXPECTED_COUNTS, EXPECTED_IDENTITIES, confirmation, dossier_bundle
from tests.fixtures.factory import SyntheticSource


def request(token, method, path, payload=None):
    con = http.client.HTTPConnection('127.0.0.1', 8765, timeout=120)
    try:
        con.request(method, path, json.dumps(payload).encode() if payload is not None else None,
                    {'Authorization': 'Bearer ' + token, 'Content-Type': 'application/json'})
        res = con.getresponse()
        return res.status, res.read(4_000_001)
    finally:
        con.close()


def json_request(token, method, path, payload=None):
    status, raw = request(token, method, path, payload)
    if status != 200:
        raise RuntimeError(f'DOSSIER_HTTP_FAILURE:{status}:{method}:{path.rsplit("/", 1)[-1]}')
    return json.loads(raw)


def main():
    if os.getenv('GITHUB_REF') not in {'refs/heads/feat/complete-dossier-journey', 'refs/heads/fix/dossier-ranges-mixed', 'refs/heads/main'} or os.getenv('GITHUB_REPOSITORY') != 'haytamAroui/caseproof':
        raise RuntimeError('protected dossier branch required')
    key = os.environ['CASEPROOF_DEV_PUBLISHABLE_KEY']
    signin = SupabasePasswordSignin(PROJECT_REF, key)
    credentials = {r: PasswordCredentials(email=os.environ[f'CASEPROOF_P3_{r.upper()}_EMAIL'],
                                          password=os.environ[f'CASEPROOF_P3_{r.upper()}_PASSWORD'])
                   for r in ('owner', 'ungranted', 'foreign')}
    sessions = {r: signin.sign_in(c) for r, c in credentials.items()}
    token = sessions['owner'].access_token
    proc = None
    try:
        proc = _server_start()
        case = success(token, 'POST', '/dev/live/matters',
                       {'org_id': os.environ['CASEPROOF_P3_ORG_ID'], 'name': 'Synthetic Belgian dossier',
                        'request_id': str(uuid4())})['case_id']
        record(os.environ['CASEPROOF_DEV_AUTH_ADMIN_KEY'], os.environ['CASEPROOF_P3_RUN_ID'], 'matter', case, 'CREATED')
        bundle = list(dossier_bundle())
        unreadable = os.getenv('CASEPROOF_DOSSIER_UNREADABLE') == 'YES'
        edge_cases = os.getenv('CASEPROOF_DOSSIER_EDGE_CASES') == 'YES'
        if edge_cases:
            from docx import Document
            from pypdf import PdfReader, PdfWriter
            draft = Document(io.BytesIO(bundle[0].content))
            draft.paragraphs[1].text = 'Voir Pièce 2, PDF p. 2.'
            draft.add_paragraph('Zie Stuk 1, PDF p. 1-99.')
            out = io.BytesIO()
            draft.save(out)
            bundle[0] = SyntheticSource(bundle[0].filename, 'draft', bundle[0].media_type, out.getvalue())
            writer = PdfWriter()
            for source in bundle[2:]:
                writer.add_page(PdfReader(io.BytesIO(source.content)).pages[0])
            mixed = io.BytesIO()
            writer.write(mixed)
            bundle[3] = SyntheticSource(bundle[3].filename, 'exhibit', 'application/pdf', mixed.getvalue())
        if unreadable:
            bundle[0] = SyntheticSource('draft.docx', 'draft', bundle[0].media_type, b'Unreadable synthetic DOCX fixture')
        status, raw = direct(token, key, '/rest/v1/rpc/caseproof_dev_register_bundle', 'POST',
                             {'p_case': case, 'p_request': str(uuid4()), 'p_sources': [s.metadata() for s in bundle]})
        if status != 200:
            raise RuntimeError('DOSSIER_REGISTRATION_FAILED')
        registered = json.loads(raw)['sources']
        versions = {s['role']: s['version_id'] for s in registered if s['role'] != 'exhibit'}
        text, scan = registered[2], registered[3]
        base = f'/v2/matters/{case}/briefs/{versions["draft"]}/dossier'
        for source, row in zip(bundle, registered, strict=True):
            path = '/'.join((os.environ['CASEPROOF_P3_ORG_ID'], case, row['document_id'], row['version_id'], source.sha256))
            record(os.environ['CASEPROOF_DEV_AUTH_ADMIN_KEY'], os.environ['CASEPROOF_P3_RUN_ID'], 'storage_file', path)
            success(token, 'POST', f'/dev/live/matters/{case}/sources/{row["document_id"]}/{row["version_id"]}/upload',
                    {'content_base64': base64.b64encode(source.content).decode()})
        for source, row in zip(bundle, registered, strict=True):
            status, original = request(token, 'GET', base + f'/sources/{row["version_id"]}/original')
            if status != 200 or hashlib.sha256(original).hexdigest() != source.sha256:
                raise RuntimeError('DOSSIER_ORIGINAL_READBACK_HASH_FAILED')
        print('DOSSIER_FOUR_SIGNED_UPLOADS_AND_ORIGINAL_HASH_READBACKS_PASSED', flush=True)
        before = json_request(token, 'GET', base)
        if before['coverage']['fully_checked'] or before['coverage']['checked'] != 0:
            raise RuntimeError('UNCONFIRMED_DOSSIER_OVERCLAIMED')
        if unreadable:
            if before['coverage']['detected'] or before['coverage']['text_coverage_complete']:
                raise RuntimeError('UNREADABLE_DOSSIER_OVERCLAIMED')
            print('DOSSIER_LIVE_UNREADABLE_DRAFT_NEVER_FULLY_CHECKED', flush=True)
            return
        success(token, 'POST', f'/dev/live/matters/{case}/sources/{scan["document_id"]}/{scan["version_id"]}/ingest?request_id={uuid4()}')
        body = confirmation(before['bundle_signature'], str(uuid4()), text['version_id'], scan['version_id'])
        if edge_cases:
            # No printed/bundle map may pre-populate the lookup. Only the
            # physical PDF reference should cause page-2 geometry retrieval.
            body['entries'][1]['pages'] = []
        json_request(token, 'POST', base + '/confirmation', body)
        report = json_request(token, 'GET', base)
        expected_identities = EXPECTED_IDENTITIES + (['resolved'] if edge_cases else [])
        expected_counts = {'detected': 9, 'checked': 3, 'unresolved': 6} if edge_cases else EXPECTED_COUNTS
        if [f['identity'] for f in report['findings']] != expected_identities:
            raise RuntimeError('DOSSIER_SEEDED_IDENTITIES_FAILED')
        counts = {k: report['coverage'][k] for k in EXPECTED_COUNTS}
        if counts != expected_counts or report['coverage']['unprocessed_regions'] < 1 or report['coverage']['fully_checked']:
            raise RuntimeError('DOSSIER_COVERAGE_OVERCLAIMED')
        if edge_cases:
            sequence = report['findings'][-1]
            mixed_source = next(s for s in report['sources'] if s['version_id'] == scan['version_id'])
            if sequence['technical_check'] != 'unresolved' or sequence['page_mapping_reason'] != 'page_sequence_requires_review':
                raise RuntimeError('DOSSIER_RANGE_FALSE_CHECK')
            if mixed_source['image_only'] or mixed_source['image_only_pages'] != [2] or report['findings'][1]['pdf_page'] != 2:
                raise RuntimeError('DOSSIER_MIXED_PDF_PAGE_STATUS_FAILED')
            print('DOSSIER_RANGE_ABSTENTION_AND_MIXED_PAGE_2_GEOMETRY_PASSED', flush=True)
        scanned = report['findings'][1]
        if scanned['ocr_uncertainty'] != 'requires_human_confirmation' or not scanned['ocr_crop']:
            raise RuntimeError('DOSSIER_OCR_UNCERTAINTY_OR_CROP_MISSING')
        for row in (text, scan):
            status, page = request(token, 'GET', base + f'/sources/{row["version_id"]}/pages/{2 if edge_cases and row == scan else 1}')
            if status != 200 or not page.startswith(b'%PDF'):
                raise RuntimeError('DOSSIER_ORIGINAL_PAGE_FAILED')
        crop = {k: v for k, v in scanned['ocr_crop'].items() if k != 'geometry_sha256'}
        status, png = request(token, 'POST', f'/dev/live/matters/{case}/sources/{scan["document_id"]}/{scan["version_id"]}/crop', crop)
        if status != 200 or not png.startswith(bytes.fromhex('89504e47')):
            raise RuntimeError('DOSSIER_CROP_FAILED')
        review = {'snapshot_hash': report['snapshot_hash'], 'reference_id': scanned['reference_id'],
                  'request_id': str(uuid4()), 'decision': 'CONFIRMED_SUPPORTED',
                  'rationale': 'Signed synthetic reviewer inspected original scan, printed label and OCR uncertainty'}
        json_request(token, 'POST', base + '/reviews', review)
        exported = json_request(token, 'GET', base + '/export')
        if len(exported['human_decisions']) != 1 or exported['human_decisions'][0]['request_id'] != review['request_id']:
            raise RuntimeError('DOSSIER_DECISION_NOT_IN_EXPORT')
        status, markdown = request(token, 'GET', base + '/export?format=markdown')
        if status != 200 or b'unresolved' not in markdown or b'CONFIRMED_SUPPORTED' not in markdown:
            raise RuntimeError('DOSSIER_MARKDOWN_EXPORT_FAILED')
        status, derivative = request(token, 'GET', base + '/export?format=docx')
        if status != 200:
            raise RuntimeError('DOSSIER_DOCX_EXPORT_FAILED')
        with ZipFile(io.BytesIO(derivative)) as archive:
            if b'scoped review annotations' not in archive.read('word/document.xml'):
                raise RuntimeError('DOSSIER_DOCX_ANNOTATIONS_MISSING')
        artifacts = Path('/tmp/caseproof-dossier-demo')
        artifacts.mkdir(exist_ok=True)
        for source in bundle:
            (artifacts / source.filename).write_bytes(source.content)
        (artifacts / 'report.json').write_text(json.dumps(exported, ensure_ascii=False, indent=2), encoding='utf-8')
        (artifacts / 'report.md').write_bytes(markdown)
        (artifacts / 'scan-crop.png').write_bytes(png)
        (artifacts / 'caseproof-annotated.docx').write_bytes(derivative)
        status, original_after_export = request(token, 'GET', base + f'/sources/{versions["draft"]}/original')
        if status != 200 or hashlib.sha256(original_after_export).hexdigest() != bundle[0].sha256:
            raise RuntimeError('DOSSIER_EXPORT_CHANGED_ORIGINAL_DRAFT')
        print('DOSSIER_PAGE_CROP_REVIEW_JSON_MARKDOWN_DOCX_PASSED', flush=True)
        _server_stop(proc)
        proc = _server_start()
        token = signin.sign_in(credentials['owner']).access_token
        reopened = json_request(token, 'GET', base + '/export')
        if reopened != exported:
            raise RuntimeError('DOSSIER_COLD_RESTART_CHANGED_SOURCES_FINDINGS_OR_DECISIONS')
        print('DOSSIER_COLD_RESTART_FRESH_SIGNIN_IDENTICAL_REPORT_PASSED', flush=True)
        # Same source set, new human mapping revision: old decision must be stale.
        body['request_id'] = str(uuid4())
        json_request(token, 'POST', base + '/confirmation', body)
        stale, _ = request(token, 'POST', base + '/reviews', review)
        current = json_request(token, 'GET', base + '/export')
        if stale != 409 or current['human_decisions'] or current['snapshot_hash'] == review['snapshot_hash']:
            raise RuntimeError('DOSSIER_STALE_REVIEW_NOT_INVALIDATED')
        print('DOSSIER_CHANGED_CONFIRMATION_INVALIDATES_OLD_REVIEW_PASSED', flush=True)
        for role in ('ungranted', 'foreign'):
            denied_token = sessions[role].access_token
            for suffix in ('', '/export', f'/sources/{scan["version_id"]}/original', f'/sources/{scan["version_id"]}/pages/1'):
                status, _ = request(denied_token, 'GET', base + suffix)
                if status not in (403, 404):
                    raise RuntimeError('DOSSIER_UNGRANTED_ACCESS_NOT_DENIED')
            status, _ = request(denied_token, 'POST', base + '/confirmation', body)
            if status not in (403, 404):
                raise RuntimeError('DOSSIER_UNGRANTED_CONFIRMATION_NOT_DENIED')
        admin_rpc(os.environ['CASEPROOF_DEV_AUTH_ADMIN_KEY'], 'revoke',
                  {'p_run': os.environ['CASEPROOF_P3_RUN_ID'], 'p_owner': str(sessions['owner'].user_id), 'p_case': case})
        for suffix in ('', '/export', f'/sources/{scan["version_id"]}/original', f'/sources/{scan["version_id"]}/pages/1'):
            status, _ = request(token, 'GET', base + suffix)
            if status not in (403, 404):
                raise RuntimeError('DOSSIER_REVOKED_ACCESS_NOT_DENIED')
        status, _ = request(token, 'POST', f'/dev/live/matters/{case}/sources/{scan["document_id"]}/{scan["version_id"]}/crop', crop)
        if status not in (403, 404):
            raise RuntimeError('DOSSIER_REVOKED_CROP_NOT_DENIED')
        print('DOSSIER_UNGRANTED_FOREIGN_REVOKED_ACCESS_DENIED', flush=True)
        print('DOSSIER_ACCEPTANCE_JSON:' + json.dumps({'commit': os.getenv('GITHUB_SHA'), **counts,
              'unprocessed_regions': report['coverage']['unprocessed_regions'],
              'fully_checked': False, 'source_count': len(registered), 'human_decisions_reopened': 1,
              'source_hashes': [s.sha256 for s in bundle], 'cleanup': 'required_by_finally_wrapper'}, sort_keys=True), flush=True)
    finally:
        _server_stop(proc)


if __name__ == '__main__':
    main()
